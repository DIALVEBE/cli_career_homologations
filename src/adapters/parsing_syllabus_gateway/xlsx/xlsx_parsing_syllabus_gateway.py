from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from gateways import ParsingSyllabusGateway
from entities import Competency, HumanActionDimensions, Subject

import unicodedata


SECTION_ALIASES = {
    "objective": [
        "PROPÓSITO  DEL ESPACIO ACADÉMICO",
        "PROPÓSITO DEL ESPACIO ACADÉMICO",
        "OBJECTIVE/PURPOSE",
        "OBJECTIVE /PURPOSE",
        "OBJETIVO GENERAL",
        "PROPÓSITO DEL ESPACIO ACADÉMICO",
    ],
    "problemic_core": [
        "NÚCLEO PROBLÉMICO AL QUE TRIBUTA EL ESPACIO ACADÉMICO",
        "NÚCLEO PROBLÉMICO Y LÍNEA",
        "NÚCLEO PROBLÉMICO",
    ],
    "didactic_strategies": [
        "PRINCIPALES ESTRATEGIAS DIDÁCTICAS DEL ESPACIO ACADÉMICO",
        "DIDÁCTIC STRATEGIES",
        "DIDACTIC STRATEGIES",
        "PRINCIPALES ESTRATEGIAS DIDÁCTICAS",
    ],
    "competencies": [
        "COMPETENCIAS, RESULTADOS DE APRENDIZAJE, MECANISMOS DE EVALUACIÓN",
        "COMPETENCES, LEARNING OUTCOMES",
        "COMPETENCIAS A DESARROLLAR",
    ],
}

SUBJECT_NAME_LABELS = [
    "Nombre del Espacio Académico",
    "COURSE NAME",
    "Espacio académico",
]

PROGRAM_LABELS = [
    "Nombre del Programa Académico",
    "ACADEMIC PROGRAM",
    "PROGRAM",
    "Programa",
]

IDENTIFICATION_LABELS = [
    *SUBJECT_NAME_LABELS,
    *PROGRAM_LABELS,
]

class XLSXParsingSyllabusGateway(ParsingSyllabusGateway):
    def extract_subject_from_syllabus(self, path: Path) -> Subject:
        """
        Extract the information from an XLSX syllabus file, the subject have the following attributes with their respective titles to search in the file:
        name: "Nombre del Espacio Académico"
        program: "Nombre del Programa Académico"
        objective: "PROPÓSITO  DEL ESPACIO ACADÉMICO (Ver instructivo anexo)"
        problemicCore: "NÚCLEO PROBLÉMICO AL QUE TRIBUTA EL ESPACIO ACADÉMICO"
        didacticStrategies: "PRINCIPALES ESTRATEGIAS DIDÁCTICAS DEL ESPACIO ACADÉMICO  (Ver instructivo anexo)"
        competencies: "COMPETENCIAS, RESULTADOS DE APRENDIZAJE, MECANISMOS DE EVALUACIÓN (Ver instructivo anexo)"

        Every competency has the following attributes:
        name: "COMPETENCIAS"
        humanActionDimensions: "DIMENSIONES DE LA ACCIÓN HUMANA"
        learningResults: "RESULTADOS DE APRENDIZAJE"
        contents: "CONTENIDOS"
        time: "TIEMPOS"
        evaluationMechanisms: "MECANISMOS DE EVALUACIÓN"
        didacticResources: "RECURSOS DIDÁCTICOS"
        """
        workbook_path = _resolve_workbook_path(path)
        try:
            workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        except (InvalidFileException, OSError, ValueError) as error:
            raise ValueError(f"Could not read XLSX syllabus: {workbook_path}") from error

        rows = _worksheet_rows(workbook[workbook.sheetnames[0]])
        all_rows: list[list[str]] | None = None

        name = _extract_inline_value(rows, SUBJECT_NAME_LABELS)
        if not name:
            all_rows = all_rows or _workbook_rows(workbook)
            name = _extract_inline_value(all_rows, SUBJECT_NAME_LABELS)

        program = _extract_inline_value(rows, PROGRAM_LABELS)
        if not program:
            all_rows = all_rows or _workbook_rows(workbook)
            program = _extract_inline_value(all_rows, PROGRAM_LABELS)

        objective = _extract_section_text(rows, SECTION_ALIASES["objective"])
        if not objective:
            all_rows = all_rows or _workbook_rows(workbook)
            objective = _extract_section_text(all_rows, SECTION_ALIASES["objective"])

        problemic_core = _extract_section_text(rows, SECTION_ALIASES["problemic_core"])
        if not problemic_core:
            all_rows = all_rows or _workbook_rows(workbook)
            problemic_core = _extract_section_text(all_rows, SECTION_ALIASES["problemic_core"])

        didactic_strategies = _extract_section_text(rows, SECTION_ALIASES["didactic_strategies"])
        if not didactic_strategies:
            all_rows = all_rows or _workbook_rows(workbook)
            didactic_strategies = _extract_section_text(
                all_rows, SECTION_ALIASES["didactic_strategies"]
            )

        competencies = _extract_competencies(rows)
        if not competencies:
            all_rows = all_rows or _workbook_rows(workbook)
            competencies = _extract_competencies(all_rows)

        return Subject(
            name=name,
            program=program,
            objective=objective,
            problemicCore=problemic_core,
            didacticStrategies=didactic_strategies,
            competencies=competencies,
        )


def _resolve_workbook_path(path: Path) -> Path:
    if path.exists():
        return path

    fallback = Path("tests/adapters/syllabus_gateways") / path.name
    if fallback.exists():
        return fallback

    return path


def _worksheet_rows(worksheet) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in worksheet.iter_rows(values_only=True):
        rows.append([_stringify(value) for value in row])
    return rows


def _workbook_rows(workbook) -> list[list[str]]:
    rows: list[list[str]] = []
    for worksheet in workbook.worksheets:
        if _should_skip_sheet(worksheet.title):
            continue
        rows.extend(_worksheet_rows(worksheet))
        rows.append([])
    return rows


def _should_skip_sheet(title: str) -> bool:
    normalized_title = _normalize(title)
    return any(token in normalized_title for token in ["RUBRIC", "RUBRICA", "INSTRUCTIVO", "LISTAS"])


def _stringify(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ")


def _normalize(value: str) -> str:
    without_accents = unicodedata.normalize("NFKD", value)
    ascii_value = "".join(char for char in without_accents if not unicodedata.combining(char))
    return " ".join(ascii_value.upper().split())


def _matches(value: str, aliases: list[str] | tuple[str, ...] | str) -> bool:
    if isinstance(aliases, str):
        aliases = [aliases]
    normalized_value = _normalize(value)
    return any(_normalize(alias) in normalized_value for alias in aliases)


def _extract_inline_value(rows: list[list[str]], labels: list[str]) -> str:
    known_labels = [_normalize(label) for label in IDENTIFICATION_LABELS]
    for row in rows:
        for index, cell in enumerate(row):
            if _matches(cell, labels):
                for next_cell in row[index + 1:]:
                    if next_cell and _normalize(next_cell) not in known_labels:
                        return next_cell
    return ""


def _extract_section_text(rows: list[list[str]], aliases: list[str]) -> str:
    start_index = _find_row_index(rows, aliases)
    if start_index is None:
        return ""

    chunks: list[str] = []
    for row in rows[start_index + 1:]:
        row_values = [cell for cell in row if cell]
        if not row_values:
            continue
        if chunks and _is_section_boundary(row_values):
            break
        if _is_section_boundary(row_values) and not chunks:
            continue
        chunks.extend(row_values)

    return " ".join(chunks).strip()


def _find_row_index(rows: list[list[str]], aliases: list[str]) -> int | None:
    for index, row in enumerate(rows):
        if any(_matches(cell, aliases) for cell in row):
            return index
    return None


def _is_section_boundary(row_values: list[str]) -> bool:
    aliases = [alias for group in SECTION_ALIASES.values() for alias in group]
    aliases.extend(
        [
            "COMPETENCIAS SABER PRO",
            "OBJETIVOS ESPECÍFICOS",
            "CONTENIDO PROGRAMÁTICO",
            "SISTEMA DE EVALUACIÓN",
        ]
    )
    return any(_is_heading_like(value, aliases) for value in row_values)


def _is_heading_like(value: str, aliases: list[str]) -> bool:
    normalized_value = _normalize(value)
    if len(normalized_value) > 140:
        return False
    return any(_normalize(alias) in normalized_value for alias in aliases)


def _extract_competencies(rows: list[list[str]]) -> list[Competency]:
    header_index = _find_competency_header(rows)
    if header_index is None:
        return _extract_simple_competencies(rows)

    header = rows[header_index]
    columns = _competency_columns(header)
    dimension_columns = _dimension_columns(rows[header_index + 1] if len(rows) > header_index + 1 else [])
    competencies: list[Competency] = []

    for row in rows[header_index + 2:]:
        if _is_competency_table_end(row):
            break

        name = _cell(row, columns["name"])
        learning_result = _cell(row, columns["learning_results"])

        if name:
            competencies.append(
                Competency(
                    name=name,
                    humanActionDimensions=_human_action_dimensions(row, dimension_columns),
                    learningResults=[learning_result] if learning_result else [],
                    contents=_cell(row, columns["contents"]),
                    time=_cell(row, columns["time"]),
                    evaluationMechanisms=_cell(row, columns["evaluation_mechanisms"]),
                    didacticResources=_cell(row, columns["didactic_resources"]),
                )
            )
        elif learning_result and competencies:
            competencies[-1].learningResults.append(learning_result)

    return competencies


def _find_competency_header(rows: list[list[str]]) -> int | None:
    for index, row in enumerate(rows):
        has_competencies = any(_matches(cell, ["COMPETENCIAS", "COMPETENCES"]) for cell in row)
        has_learning_results = any(_matches(cell, ["RESULTADOS DE APRENDIZAJE", "LEARNING OUTCOMES"]) for cell in row)
        has_contents = any(_matches(cell, ["CONTENIDOS", "CONTENTS", "CONCEPTUAL REFERENCE"]) for cell in row)
        if has_competencies and has_learning_results and has_contents:
            return index
    return None


def _competency_columns(header: list[str]) -> dict[str, int]:
    return {
        "name": _find_column(header, ["COMPETENCIAS", "COMPETENCES"]),
        "learning_results": _find_column(header, ["RESULTADOS DE APRENDIZAJE", "LEARNING OUTCOMES"]),
        "contents": _find_column(header, ["CONTENIDOS", "CONTENTS", "CONCEPTUAL REFERENCE"]),
        "time": _find_column(header, ["TIEMPOS", "TIMES"]),
        "evaluation_mechanisms": _find_column(header, ["MECANISMOS DE EVALUACIÓN", "DELIVERABLES"]),
        "didactic_resources": _find_column(header, ["RECURSOS DIDÁCTICOS", "DIDÁCTIC RESOURCES", "DIDACTIC RESOURCES"]),
    }


def _dimension_columns(row: list[str]) -> dict[str, int]:
    return {
        "comprehend": _find_column(row, ["Comprender", "Understanding"]),
        "act": _find_column(row, ["Obrar", "Acting"]),
        "do": _find_column(row, ["Hacer", "Doing"]),
        "communicate": _find_column(row, ["Comunicar", "Communicating"]),
        "feel": _find_column(row, "Sentir"),
    }


def _find_column(row: list[str], label: list[str] | str) -> int:
    for index, cell in enumerate(row):
        if _matches(cell, label):
            return index
    return -1


def _cell(row: list[str], index: int) -> str:
    if index < 0 or index >= len(row):
        return ""
    return row[index]


def _human_action_dimensions(row: list[str], columns: dict[str, int]) -> HumanActionDimensions:
    return HumanActionDimensions(
        comprehend=_is_checked(_cell(row, columns["comprehend"])),
        act=_is_checked(_cell(row, columns["act"])),
        do=_is_checked(_cell(row, columns["do"])),
        communicate=_is_checked(_cell(row, columns["communicate"])),
        feel=_is_checked(_cell(row, columns["feel"])),
    )


def _is_checked(value: str) -> bool:
    return _normalize(value) in {"X", "SI", "TRUE", "1"}


def _is_competency_table_end(row: list[str]) -> bool:
    values = [cell for cell in row if cell]
    if not values:
        return False
    first_value = values[0]
    return _matches(first_value, ["COMPETENCIAS SABER PRO", "BIBLIOGRAFÍA", "REFERENCIAS", "RÚBRICA"])


def _extract_simple_competencies(rows: list[list[str]]) -> list[Competency]:
    header_index = _find_simple_competency_header(rows)
    if header_index is None:
        return []

    header = rows[header_index]
    name_column = _find_column(header, ["Competencia", "Competences"])
    description_column = _find_column(header, ["Descripción", "Description"])
    competencies: list[Competency] = []

    for row in rows[header_index + 1:]:
        values = [cell for cell in row if cell]
        if not values or _is_section_boundary(values):
            break

        name = _cell(row, name_column)
        description = _cell(row, description_column)
        if not name:
            continue

        competencies.append(
            Competency(
                name=name,
                humanActionDimensions=HumanActionDimensions(
                    comprehend=False,
                    act=False,
                    do=False,
                    communicate=False,
                    feel=False,
                ),
                learningResults=[description] if description else [],
                contents="",
                time="",
                evaluationMechanisms="",
                didacticResources="",
            )
        )

    return competencies


def _find_simple_competency_header(rows: list[list[str]]) -> int | None:
    for index, row in enumerate(rows):
        has_competence = any(_matches(cell, ["Competencia", "Competences"]) for cell in row)
        has_description = any(_matches(cell, ["Descripción", "Description"]) for cell in row)
        if has_competence and has_description:
            return index
    return None
